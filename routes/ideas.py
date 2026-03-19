"""
Ideas routes blueprint: CRUD, upload, transcription, analysis, bulk operations, CSV export.
"""

import os
import io
import csv
import json
import uuid
import tempfile
import threading
from datetime import datetime

from flask import Blueprint, request, jsonify, session, Response, send_file, current_app
from werkzeug.utils import secure_filename

from config import (
    ALLOWED_AUDIO_EXTENSIONS, ALLOWED_DOCUMENT_EXTENSIONS,
    upload_jobs, logger
)
from database import get_db
from auth import login_required, reviewer_required, admin_required
from services.backup import save_ideas_backup
from services.transcription import (
    _transcribe_with_elevenlabs, _split_audio_chunks,
    _clean_hallucinations, _is_whisper_hallucination,
    _process_upload
)
from services.analysis import _auto_analyze, _get_company_context_for_prompt
from utils.validation import validate_limit, validate_offset

ideas_bp = Blueprint('ideas', __name__)


@ideas_bp.route('/api/ideas', methods=['GET'])
@login_required
def api_ideas():
    db = get_db()
    filters = []
    params = []

    dept = request.args.get('department')
    role = request.args.get('role')
    status = request.args.get('status')
    search = request.args.get('search')
    limit = validate_limit(request.args.get('limit'), default=50, max_val=200)
    offset = validate_offset(request.args.get('offset'))

    # Submitters only see their own ideas or company-wide ones
    user_role = session.get('user_role')
    if user_role == 'submitter':
        filters.append("(author_id = ? OR visibility = 'company')")
        params.append(session['user_id'])

    if dept:
        filters.append('department = ?')
        params.append(dept)
    if role:
        filters.append('role = ?')
        params.append(role)
    if status:
        filters.append('status = ?')
        params.append(status)
    if search:
        filters.append('(transcript LIKE ? OR author_name LIKE ?)')
        params.extend([f'%{search}%', f'%{search}%'])

    idea_type = request.args.get('idea_type')
    if idea_type:
        filters.append('idea_type = ?')
        params.append(idea_type)

    where = ('WHERE ' + ' AND '.join(filters)) if filters else ''
    total = db.execute(f'SELECT COUNT(*) FROM ideas {where}', params).fetchone()[0]
    listing_cols = ('id, author_id, author_name, department, role, audio_filename, '
                    'duration_seconds, transcript, status, ai_score, ai_analysis, '
                    'reviewer_note, reviewed_by, reviewed_at, created_at, '
                    'visibility, tags, assigned_to, deadline, campaign_id, '
                    'transcribed_at, stt_engine, idea_type')
    try:
        rows = db.execute(f'SELECT {listing_cols} FROM ideas {where} ORDER BY created_at DESC LIMIT ? OFFSET ?',
                          params + [limit, offset]).fetchall()
    except Exception as e:
        logger.warning('Listing query fallback due to: %s', e)
        rows = db.execute(f'SELECT * FROM ideas {where} ORDER BY created_at DESC LIMIT ? OFFSET ?',
                          params + [limit, offset]).fetchall()
    db.close()
    data = []
    for r in rows:
        d = dict(r)
        d.pop('audio_data', None)
        d.pop('ai_analysis', None)
        d['has_audio'] = bool(d.get('audio_filename'))
        data.append(d)
    return jsonify({'data': data, 'total': total})


@ideas_bp.route('/api/ideas/<int:idea_id>', methods=['GET'])
@login_required
def api_idea_detail(idea_id):
    db = get_db()
    idea = db.execute('SELECT * FROM ideas WHERE id = ?', (idea_id,)).fetchone()
    db.close()
    if not idea:
        return jsonify({'error': 'Napad nenajdeny', 'code': 'not_found'}), 404
    user_role = session.get('user_role')
    if user_role == 'submitter':
        if idea['author_id'] != session['user_id'] and idea['visibility'] != 'company':
            return jsonify({'error': 'Pristup zamietnuty', 'code': 'forbidden'}), 403
    return jsonify(dict(idea))


@ideas_bp.route('/api/ideas/<int:idea_id>', methods=['PATCH'])
@reviewer_required
def api_idea_update(idea_id):
    data = request.get_json() or {}
    allowed = {'status', 'reviewer_note', 'visibility', 'tags', 'assigned_to', 'deadline', 'idea_type'}
    updates = {k: v for k, v in data.items() if k in allowed}
    if 'visibility' in updates and updates['visibility'] not in ('personal', 'company'):
        return jsonify({'error': 'Neplatna hodnota viditelnosti', 'code': 'invalid_visibility'}), 400
    if 'status' in updates and updates['status'] not in ('new', 'in_review', 'accepted', 'rejected', 'v_realizacii'):
        return jsonify({'error': 'Neplatny status', 'code': 'invalid_status'}), 400
    if 'idea_type' in updates and updates['idea_type'] not in ('napad', 'porada'):
        return jsonify({'error': 'Neplatny typ zaznamu', 'code': 'invalid_type'}), 400
    if 'tags' in updates:
        try:
            parsed = json.loads(updates['tags']) if isinstance(updates['tags'], str) else updates['tags']
            updates['tags'] = json.dumps([str(t) for t in parsed[:10]], ensure_ascii=False)
        except Exception:
            return jsonify({'error': 'Neplatny format tagov', 'code': 'invalid_tags'}), 400
    if not updates:
        return jsonify({'error': 'Nic na aktualizaciu', 'code': 'no_updates'}), 400

    db = get_db()
    idea = db.execute('SELECT * FROM ideas WHERE id = ?', (idea_id,)).fetchone()
    if not idea:
        db.close()
        return jsonify({'error': 'Napad nenajdeny', 'code': 'not_found'}), 404

    if 'status' in updates:
        updates['reviewed_by'] = session['user_name']
        updates['reviewed_at'] = datetime.now().isoformat()

    set_clause = ', '.join(f'{k} = ?' for k in updates)
    values = list(updates.values()) + [idea_id]
    db.execute(f'UPDATE ideas SET {set_clause} WHERE id = ?', values)
    db.commit()
    db.close()
    save_ideas_backup()
    return jsonify({'ok': True})


@ideas_bp.route('/api/ideas/<int:idea_id>', methods=['DELETE'])
@admin_required
def api_idea_delete(idea_id):
    try:
        db = get_db()
        db.execute('DELETE FROM comments WHERE idea_id = ?', (idea_id,))
        db.execute('DELETE FROM votes WHERE idea_id = ?', (idea_id,))
        db.execute('DELETE FROM meeting_ideas WHERE idea_id = ?', (idea_id,))
        db.execute('DELETE FROM ideas WHERE id = ?', (idea_id,))
        db.commit()
        db.close()
        save_ideas_backup()
        return jsonify({'ok': True})
    except Exception as e:
        logger.error('Delete idea %s error: %s', idea_id, e)
        return jsonify({'error': f'Chyba pri mazani: {str(e)[:200]}', 'code': 'delete_error'}), 500


@ideas_bp.route('/api/ideas/upload', methods=['POST'])
@login_required
def api_ideas_upload():
    logger.info('Upload request from user_id=%s, origin=%s', session.get('user_id'), request.headers.get('Origin', '?'))
    api_key = os.environ.get('OPENAI_API_KEY', '')
    el_key = os.environ.get('ELEVENLABS_API_KEY', '')
    if not api_key and not el_key:
        return jsonify({'error': 'Ani ElevenLabs ani OpenAI API kluc nie je nastaveny', 'code': 'no_api_key'}), 500

    if 'audio' not in request.files:
        return jsonify({'error': 'Chyba audio subor', 'code': 'missing_file'}), 400

    file = request.files['audio']
    if not file.filename:
        return jsonify({'error': 'Prazdny subor', 'code': 'empty_file'}), 400

    ext = os.path.splitext(secure_filename(file.filename))[1].lower()
    if ext not in ALLOWED_AUDIO_EXTENSIONS:
        return jsonify({'error': f'Nepodporovany format: {ext}', 'code': 'unsupported_format'}), 400

    department = (request.form.get('department') or '').strip()
    role = (request.form.get('role') or '').strip()
    visibility = (request.form.get('visibility') or 'personal').strip()
    if visibility not in ('personal', 'company'):
        visibility = 'personal'

    if not department or not role:
        return jsonify({'error': 'Oddelenie a rola su povinne', 'code': 'missing_fields'}), 400

    try:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            file.save(tmp)
            tmp_path = tmp.name
    except Exception as e:
        return jsonify({'error': f'Chyba pri ukladani: {str(e)}', 'code': 'save_error'}), 500

    job_id = str(uuid.uuid4())
    user_id = session['user_id']
    user_name = session['user_name']
    upload_jobs.set(job_id, {'status': 'processing'})

    # Capture app reference for the background thread
    app = current_app._get_current_object()
    t = threading.Thread(
        target=_process_upload,
        args=(app, job_id, tmp_path, ext, user_id, user_name, department, role, visibility, api_key),
        daemon=True
    )
    t.start()

    return jsonify({'job_id': job_id, 'status': 'processing'})


@ideas_bp.route('/api/ideas/job/<job_id>', methods=['GET'])
@login_required
def api_ideas_job(job_id):
    import time
    job = upload_jobs.get(job_id)
    logger.debug('Job poll: %s, found=%s, status=%s, total_jobs=%d', job_id, job is not None, job['status'] if job else 'N/A', len(upload_jobs))
    if not job:
        return jsonify({'error': 'Job nenajdeny', 'code': 'not_found'}), 404
    if job['status'] == 'done':
        result = job['result']
        if 'completed_at' not in job:
            upload_jobs.update(job_id, 'completed_at', time.time())
        elif time.time() - job['completed_at'] > 60:
            upload_jobs.delete(job_id)
        return jsonify(result)
    elif job['status'] == 'error':
        err = job.get('error', 'Neznama chyba')
        if 'completed_at' not in job:
            upload_jobs.update(job_id, 'completed_at', time.time())
        elif time.time() - job['completed_at'] > 60:
            upload_jobs.delete(job_id)
        return jsonify({'error': err, 'code': 'job_error'}), 500
    else:
        return jsonify({'status': 'processing'}), 202


@ideas_bp.route('/api/ideas/my-recent', methods=['GET'])
@login_required
def api_my_recent_ideas():
    """Return last 10 ideas for the current user (for recorder feedback)."""
    user_id = session['user_id']
    db = get_db()
    rows = db.execute('''
        SELECT id, transcript, department, status, visibility, created_at, duration_seconds, stt_engine
        FROM ideas WHERE author_id = ? ORDER BY id DESC LIMIT 10
    ''', (user_id,)).fetchall()
    db.close()
    ideas = []
    for r in rows:
        t = r['transcript'] or ''
        ideas.append({
            'id': r['id'],
            'preview': (t[:80] + '...') if len(t) > 80 else t,
            'department': r['department'],
            'status': r['status'],
            'visibility': r['visibility'] or 'personal',
            'created_at': r['created_at'],
            'duration': r['duration_seconds'] or 0,
            'stt_engine': r['stt_engine'] or ''
        })
    return jsonify(ideas)


@ideas_bp.route('/api/ideas/<int:idea_id>/audio', methods=['GET'])
@login_required
def api_idea_audio(idea_id):
    import base64
    db = get_db()
    idea = db.execute('SELECT audio_filename, audio_data FROM ideas WHERE id = ?', (idea_id,)).fetchone()
    db.close()
    if not idea or not idea['audio_filename']:
        return jsonify({'error': 'Audio nenajdene', 'code': 'not_found'}), 404
    # Try file on disk first
    audio_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'audio_uploads')
    audio_path = os.path.join(audio_dir, idea['audio_filename'])
    if os.path.exists(audio_path):
        return send_file(audio_path, as_attachment=False)
    # Fallback: serve from DB (base64)
    audio_b64 = idea['audio_data'] if idea['audio_data'] else ''
    if not audio_b64:
        return jsonify({'error': 'Subor neexistuje', 'code': 'file_missing'}), 404
    audio_bytes = base64.b64decode(audio_b64)
    fname = idea['audio_filename']
    mime = 'audio/webm' if fname.endswith('.webm') else 'audio/mpeg' if fname.endswith('.mp3') else 'audio/ogg' if fname.endswith('.ogg') else 'audio/wav'
    return Response(audio_bytes, mimetype=mime, headers={'Content-Disposition': f'inline; filename="{fname}"'})


@ideas_bp.route('/api/ideas/<int:idea_id>/retranscribe', methods=['POST'])
@login_required
def api_idea_retranscribe(idea_id):
    """Re-run transcription on stored audio_data."""
    import base64
    import openai

    db = get_db()
    idea = db.execute('SELECT audio_data, audio_filename, duration_seconds FROM ideas WHERE id = ?', (idea_id,)).fetchone()
    db.close()
    if not idea or not idea['audio_data']:
        return jsonify({'error': 'Audio nie je k dispozicii pre tento napad', 'code': 'no_audio'}), 404
    audio_bytes = base64.b64decode(idea['audio_data'])
    ext = os.path.splitext(idea['audio_filename'])[1] if idea['audio_filename'] else '.webm'
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(audio_bytes)
        tmp_path = tmp.name
    try:
        transcript_text = ''
        total_duration = 0
        stt_warning = None

        # PRIMARY: Try ElevenLabs Scribe
        el_text, el_duration, el_warning = _transcribe_with_elevenlabs(tmp_path)
        if el_warning:
            stt_warning = el_warning
        if el_text:
            transcript_text = el_text
            total_duration = el_duration
        elif el_text is not None:
            transcript_text = ''
            total_duration = el_duration
        else:
            # FALLBACK: Whisper
            api_key = os.environ.get('OPENAI_API_KEY')
            if not api_key:
                return jsonify({'error': 'Ani ElevenLabs ani OpenAI API kluc nie je nastaveny', 'code': 'no_api_key'}), 500
            client = openai.OpenAI(api_key=api_key)
            chunks = _split_audio_chunks(tmp_path)
            all_text = []
            for i, chunk_path in enumerate(chunks):
                if os.path.getsize(chunk_path) > 25 * 1024 * 1024:
                    continue
                with open(chunk_path, 'rb') as f:
                    tr = client.audio.transcriptions.create(
                        model='whisper-1', file=f, language='sk',
                        response_format='verbose_json',
                        prompt='Toto je nahravka napadu alebo myslienky v slovencine.' if i == 0 else (all_text[-1][-200:] if all_text else '')
                    )
                cleaned = _clean_hallucinations(tr.text or '')
                total_duration += int(getattr(tr, 'duration', 0) or 0)
                if cleaned:
                    all_text.append(cleaned)
            for cp in chunks:
                if cp != tmp_path and os.path.exists(cp):
                    try:
                        os.unlink(cp)
                    except Exception as e:
                        logger.warning('Failed to clean up chunk: %s', e)
            transcript_text = ' '.join(all_text).strip()

        if transcript_text:
            transcript_text = _clean_hallucinations(transcript_text)
        transcript_text = transcript_text or '[Transkript nedostupny]'
        retranscribe_engine = 'elevenlabs' if el_text is not None else 'whisper'
        retranscribe_time = datetime.now().isoformat()

        db2 = get_db()
        db2.execute('UPDATE ideas SET transcript = ?, duration_seconds = ?, transcribed_at = ?, stt_engine = ? WHERE id = ?',
                    (transcript_text, total_duration or idea['duration_seconds'] if 'duration_seconds' in idea.keys() else 0, retranscribe_time, retranscribe_engine, idea_id))
        db2.commit()
        db2.close()
        save_ideas_backup()
        result = {'ok': True, 'transcript': transcript_text, 'transcribed_at': retranscribe_time, 'stt_engine': retranscribe_engine}
        if stt_warning:
            result['warning'] = stt_warning
        return jsonify(result)
    except Exception as e:
        logger.error('Retranscribe error for idea %s: %s', idea_id, e)
        return jsonify({'error': str(e), 'code': 'retranscribe_error'}), 500
    finally:
        try:
            os.unlink(tmp_path)
        except Exception as e:
            logger.warning('Failed to clean up tmp file: %s', e)


@ideas_bp.route('/api/ideas/<int:idea_id>/analyze', methods=['POST'])
@login_required
def api_idea_analyze(idea_id):
    import anthropic as anthropic_sdk

    api_key = os.environ.get('ANTHROPIC_API_KEY')
    if not api_key:
        return jsonify({'error': 'Anthropic API kluc nie je nastaveny', 'code': 'no_api_key'}), 500

    db = get_db()
    idea = db.execute('SELECT * FROM ideas WHERE id = ?', (idea_id,)).fetchone()
    if not idea:
        db.close()
        return jsonify({'error': 'Napad nenajdeny', 'code': 'not_found'}), 404

    transcript = idea['transcript']
    if not transcript:
        db.close()
        return jsonify({'error': 'Chyba transkript', 'code': 'no_transcript'}), 400

    company_context = _get_company_context_for_prompt()

    prompt = f"""Analyzuj nasledujuci interny napad od zamestnanca a ohodno ho.

{('--- KONTEXT FIRMY ---' + chr(10) + company_context + chr(10) + '--- KONIEC KONTEXTU ---' + chr(10)) if company_context else ''}
Oddelenie: {idea['department']}
Rola: {idea['role']}
Transkript napadu:
"{transcript}"

Vrat JSON s tymto formatom (iba JSON, bez markdown):
{{
  "score": <1-10>,
  "clarity": <1-10>,
  "feasibility": <1-10>,
  "relevance": <1-10>,
  "summary": "<2-3 vety zhrnutie napadu>",
  "strengths": ["<silna stranka 1>", "<silna stranka 2>"],
  "weaknesses": ["<slaba stranka 1>"],
  "next_steps": ["<konkretny krok 1>", "<konkretny krok 2>"],
  "category": "<one of: process_improvement|cost_reduction|revenue|product|other>",
  "tags": ["<tag1>", "<tag2>"]
}}

Hodnot objektivne. score je celkove hodnotenie potencialu napadu.
relevance je hodnotenie relevancie napadu pre firmu (ak je k dispozicii kontext firmy, zohladni ciele, priority a hodnoty firmy).
Pre tags pouzi max 5 tagov z tohto zoznamu (alebo vlastne slovenske/anglicke slovo): quick_win, cost_reduction, product, process, customer, technical, innovation, urgent, automation, hr, marketing, quality."""

    try:
        client = anthropic_sdk.Anthropic(api_key=api_key)
        message = client.messages.create(
            model='claude-haiku-4-5-20251001',
            max_tokens=1000,
            messages=[{'role': 'user', 'content': prompt}]
        )
        raw = message.content[0].text.strip()
        if raw.startswith('```'):
            raw = raw.split('```')[1]
            if raw.startswith('json'):
                raw = raw[4:]
        analysis = json.loads(raw)
        score = int(analysis.get('score', 0))
        tags = json.dumps(analysis.get('tags', []), ensure_ascii=False)

        db.execute('UPDATE ideas SET ai_score = ?, ai_analysis = ?, tags = ? WHERE id = ?',
                   (score, json.dumps(analysis, ensure_ascii=False), tags, idea_id))
        db.commit()
        db.close()
        save_ideas_backup()
        return jsonify({'ok': True, 'analysis': analysis})
    except Exception as e:
        db.close()
        logger.error('Analyze error: %s', e)
        return jsonify({'error': f'AI analyza zlyhala: {str(e)}', 'code': 'analyze_error'}), 500


@ideas_bp.route('/api/ideas/text', methods=['POST'])
@login_required
def api_ideas_text():
    """Create an idea from text input."""
    data = request.get_json() or {}
    text = (data.get('text') or '').strip()
    department = (data.get('department') or '').strip()
    role = (data.get('role') or '').strip()
    visibility = (data.get('visibility') or 'personal').strip()
    if visibility not in ('personal', 'company'):
        visibility = 'personal'
    if not text:
        return jsonify({'error': 'Text napadu je povinny', 'code': 'missing_text'}), 400
    if not department or not role:
        return jsonify({'error': 'Oddelenie a rola su povinne', 'code': 'missing_fields'}), 400

    db = get_db()
    cursor = db.execute('''
        INSERT INTO ideas (author_id, author_name, department, role, duration_seconds, transcript, status, visibility)
        VALUES (?, ?, ?, ?, 0, ?, 'new', ?)
    ''', (session['user_id'], session['user_name'], department, role, text, visibility))
    idea_id = cursor.lastrowid
    db.commit()
    db.close()
    save_ideas_backup()

    # Auto-analyze in background
    app = current_app._get_current_object()
    threading.Thread(target=_auto_analyze, args=(app, idea_id,), daemon=True).start()

    return jsonify({'ok': True, 'id': idea_id, 'message': 'Napad uspesne vytvoreny'}), 201


@ideas_bp.route('/api/ideas/upload-document', methods=['POST'])
@login_required
def api_ideas_upload_document():
    """Create an idea from an uploaded document (PDF, DOCX, TXT, image)."""
    if 'document' not in request.files:
        return jsonify({'error': 'Chyba subor', 'code': 'missing_file'}), 400

    file = request.files['document']
    if not file.filename:
        return jsonify({'error': 'Prazdny subor', 'code': 'empty_file'}), 400

    ext = os.path.splitext(secure_filename(file.filename))[1].lower()
    if ext not in ALLOWED_DOCUMENT_EXTENSIONS:
        return jsonify({'error': f'Nepodporovany format: {ext}. Podporovane: PDF, DOCX, TXT, MD, obrazky', 'code': 'unsupported_format'}), 400

    department = (request.form.get('department') or '').strip()
    role = (request.form.get('role') or '').strip()
    visibility = (request.form.get('visibility') or 'personal').strip()
    if visibility not in ('personal', 'company'):
        visibility = 'personal'
    if not department or not role:
        return jsonify({'error': 'Oddelenie a rola su povinne', 'code': 'missing_fields'}), 400

    try:
        content = file.read()
        text = ''

        if ext in ('.txt', '.md', '.rtf'):
            text = content.decode('utf-8', errors='replace')
        elif ext == '.pdf':
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(content))
                pages = []
                for page in reader.pages:
                    pages.append(page.extract_text() or '')
                text = '\n'.join(pages)
            except ImportError:
                try:
                    from pdfminer.high_level import extract_text as pdf_extract
                    text = pdf_extract(io.BytesIO(content))
                except ImportError:
                    text = f'[PDF subor: {file.filename} - kniznica na citanie PDF nie je nainstalovana]'
        elif ext in ('.docx', '.doc'):
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                text = '\n'.join([p.text for p in doc.paragraphs])
            except ImportError:
                text = f'[DOCX subor: {file.filename} - kniznica na citanie DOCX nie je nainstalovana]'
        elif ext in ('.png', '.jpg', '.jpeg', '.gif', '.webp'):
            text = f'[Obrazok: {file.filename}]'
            try:
                import pytesseract
                from PIL import Image
                img = Image.open(io.BytesIO(content))
                ocr_text = pytesseract.image_to_string(img, lang='slk+eng')
                if ocr_text.strip():
                    text = f'[Obrazok: {file.filename}]\n\n{ocr_text.strip()}'
            except ImportError:
                pass
            except Exception as ocr_err:
                logger.warning('OCR error: %s', ocr_err)

        if not text.strip():
            text = f'[Importovany subor: {file.filename}]'

        if len(text) > 50000:
            text = text[:50000] + '\n\n[... text skrateny, povodny subor mal viac ako 50000 znakov]'

        db = get_db()
        cursor = db.execute('''
            INSERT INTO ideas (author_id, author_name, department, role, duration_seconds, transcript, status, visibility)
            VALUES (?, ?, ?, ?, 0, ?, 'new', ?)
        ''', (session['user_id'], session['user_name'], department, role, text, visibility))
        idea_id = cursor.lastrowid
        db.commit()
        db.close()
        save_ideas_backup()

        app = current_app._get_current_object()
        threading.Thread(target=_auto_analyze, args=(app, idea_id,), daemon=True).start()

        return jsonify({'ok': True, 'id': idea_id, 'message': 'Dokument uspesne importovany ako napad', 'transcript': text[:200]}), 201

    except Exception as e:
        logger.error('Document upload error: %s', e)
        return jsonify({'error': f'Chyba pri spracovani suboru: {str(e)}', 'code': 'upload_error'}), 500


@ideas_bp.route('/api/ideas/bulk-delete', methods=['POST'])
@admin_required
def api_ideas_bulk_delete():
    data = request.get_json() or {}
    ids = data.get('ids', [])
    if not ids or not isinstance(ids, list):
        return jsonify({'error': 'Ziadne napady na vymazanie', 'code': 'no_ids'}), 400
    try:
        db = get_db()
        placeholders = ','.join(['?'] * len(ids))
        db.execute(f'DELETE FROM comments WHERE idea_id IN ({placeholders})', ids)
        db.execute(f'DELETE FROM votes WHERE idea_id IN ({placeholders})', ids)
        db.execute(f'DELETE FROM meeting_ideas WHERE idea_id IN ({placeholders})', ids)
        db.execute(f'DELETE FROM ideas WHERE id IN ({placeholders})', ids)
        db.commit()
        db.close()
        save_ideas_backup()
        return jsonify({'ok': True, 'deleted': len(ids)})
    except Exception as e:
        logger.error('Bulk delete error: %s', e)
        return jsonify({'error': f'Chyba pri mazani: {str(e)[:200]}', 'code': 'bulk_delete_error'}), 500


@ideas_bp.route('/api/ideas/bulk-update', methods=['POST'])
@reviewer_required
def api_ideas_bulk_update():
    data = request.get_json() or {}
    ids = data.get('ids', [])
    updates = data.get('updates', {})
    if not ids or not isinstance(ids, list):
        return jsonify({'error': 'Ziadne zaznamy na aktualizaciu', 'code': 'no_ids'}), 400
    allowed = {'idea_type', 'status', 'visibility'}
    updates = {k: v for k, v in updates.items() if k in allowed}
    if not updates:
        return jsonify({'error': 'Nic na aktualizaciu', 'code': 'no_updates'}), 400
    if 'idea_type' in updates and updates['idea_type'] not in ('napad', 'porada'):
        return jsonify({'error': 'Neplatny typ zaznamu', 'code': 'invalid_type'}), 400
    db = get_db()
    placeholders = ','.join(['?'] * len(ids))
    set_clause = ', '.join(f'{k} = ?' for k in updates)
    values = list(updates.values()) + ids
    db.execute(f'UPDATE ideas SET {set_clause} WHERE id IN ({placeholders})', values)
    db.commit()
    db.close()
    save_ideas_backup()
    return jsonify({'ok': True, 'updated': len(ids)})


@ideas_bp.route('/api/ideas/export-csv')
@login_required
def api_ideas_export_csv():
    db = get_db()
    filters = []
    params = []

    dept = request.args.get('department')
    role = request.args.get('role')
    status = request.args.get('status')
    search = request.args.get('search')

    user_role = session.get('user_role')
    if user_role == 'submitter':
        filters.append("(author_id = ? OR visibility = 'company')")
        params.append(session['user_id'])

    if dept:
        filters.append('department = ?')
        params.append(dept)
    if role:
        filters.append('role = ?')
        params.append(role)
    if status:
        filters.append('status = ?')
        params.append(status)
    if search:
        filters.append('(transcript LIKE ? OR author_name LIKE ?)')
        params.extend([f'%{search}%', f'%{search}%'])

    where = ('WHERE ' + ' AND '.join(filters)) if filters else ''
    export_cols = ('id, author_name, department, role, transcript, ai_score, status, '
                   'visibility, assigned_to, deadline, tags, created_at, idea_type')
    try:
        rows = db.execute(f'SELECT {export_cols} FROM ideas {where} ORDER BY created_at DESC', params).fetchall()
    except Exception as e:
        logger.warning('CSV export fallback due to: %s', e)
        rows = db.execute(f'SELECT * FROM ideas {where} ORDER BY created_at DESC', params).fetchall()
    db.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Autor', 'Oddelenie', 'Rola', 'Transkript', 'AI Skore', 'Status', 'Viditelnost', 'Priradene', 'Deadline', 'Tagy', 'Vytvorene'])
    for r in rows:
        d = dict(r)
        writer.writerow([
            d.get('id', ''),
            d.get('author_name', ''),
            d.get('department', ''),
            d.get('role', ''),
            d.get('transcript', ''),
            d.get('ai_score', ''),
            d.get('status', ''),
            d.get('visibility', ''),
            d.get('assigned_to', ''),
            d.get('deadline', ''),
            d.get('tags', ''),
            d.get('created_at', '')
        ])

    csv_data = output.getvalue()
    return Response(
        csv_data,
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename=ridea-napady-{datetime.now().strftime("%Y%m%d")}.csv'}
    )
